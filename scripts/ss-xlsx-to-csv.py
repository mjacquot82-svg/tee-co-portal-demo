#!/usr/bin/env python3
"""Convert the S&S Data Library Excel sheets (Products.xlsx, Specs.xlsx) to CSV for import-ss-products.mjs.
Usage: python3 scripts/ss-xlsx-to-csv.py <folder with the .xlsx files> <output folder>
Needs openpyxl (pip install openpyxl). Reads only; never touches any database."""
import csv, os, sys
from openpyxl import load_workbook

def convert(src, dst):
    wb = load_workbook(src, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    with open(dst, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        for row in ws.iter_rows(values_only=True):
            w.writerow(["" if v is None else v for v in row])

if __name__ == "__main__":
    src_dir, out_dir = sys.argv[1], sys.argv[2]
    os.makedirs(out_dir, exist_ok=True)
    for name in ("Products", "Specs"):
        src = os.path.join(src_dir, f"{name}.xlsx")
        if os.path.exists(src):
            convert(src, os.path.join(out_dir, f"{name}.csv"))
            print(f"converted {name}.xlsx")
